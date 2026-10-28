"""Chek va qarz hisobotini Excel qilish.

ChekHisobot'dagi asosiy qulaylik ko'chirildi: chek qatori yonidagi «+» ni
bosganda o'sha chekning TOVARLARI ochiladi (qatorlarni guruhlash / outline).
Shunda bir ekranda o'nlab chek ko'rinadi, kerakligini ochib ko'rasiz.

DIQQAT (ChekHisobot izohidan): telefondagi Excel va Telegram ko'ruvchisi
guruhlash tugmalarini ko'rsatmaydi — shuning uchun tovarlar sukut bo'yicha
OCHIQ turadi. Kompyuterda chap chetdagi «1» bilan yig'ib qo'ysa bo'ladi.
"""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database.channel_db import DB_PATH

_OUT_DIR = os.path.join(os.path.dirname(DB_PATH), "exports")
os.makedirs(_OUT_DIR, exist_ok=True)

_HEAD_FILL = PatternFill("solid", fgColor="1F6F43")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_CHEK_FILL = PatternFill("solid", fgColor="EAF3EC")
_MONEY = "#,##0"

_COLS = ["Sana", "Vaqt", "Chek", "Tovar / izoh", "Miqdor", "Narx",
         "Summa", "Qarzga", "Qoldiq"]


def _safe(name: str) -> str:
    """Fayl nomiga yaroqli ko'rinish."""
    keep = [c if (c.isalnum() or c in " -_") else "_" for c in (name or "")]
    return "".join(keep).strip().replace(" ", "_")[:40] or "mijoz"


def _sheet(wb: Workbook, title: str):
    ws = wb.active
    ws.title = (title or "Cheklar")[:31]
    # Chek qatori guruhning TEPASIDA turadi, tovarlar ostida — shuning uchun
    # «summary below» o'chiriladi, aks holda Excel «+» ni oxirgi tovardan
    # keyin chizadi va qaysi chekniki ekani ko'rinmaydi.
    ws.sheet_properties.outlinePr.summaryBelow = False
    for i, h in enumerate(_COLS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font = _HEAD_FILL, _HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for i, w in enumerate([11, 7, 9, 44, 9, 12, 13, 12, 13], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _money_cell(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=round(float(value or 0)))
    c.number_format = _MONEY
    return c


def receipts_xlsx(rows, title: str, path: str = "") -> str:
    """`rows` — har biri {chek, tovarlar} lug'ati. Faylning yo'lini qaytaradi.

    Kutilgan shakl:
        {"at", "doc_no", "kind", "payable", "debt", "client_debt",
         "client", "items": [{"name","qty","price","amount"}]}
    """
    wb = Workbook()
    ws = _sheet(wb, title or "Cheklar")

    r = 2
    jami = qarzga = 0.0
    for rec in rows:
        at = (rec.get("at") or "")
        sana, vaqt = (at.split(" ") + [""])[:2]
        ws.cell(row=r, column=1, value=sana)
        ws.cell(row=r, column=2, value=vaqt[:5])
        ws.cell(row=r, column=3, value=rec.get("doc_no") or "")
        izoh = rec.get("kind") or ""
        if rec.get("client"):
            izoh = f"{izoh} · {rec['client']}"
        ws.cell(row=r, column=4, value=izoh)
        _money_cell(ws, r, 7, rec.get("payable"))
        _money_cell(ws, r, 8, rec.get("debt"))
        if rec.get("client_debt") is not None:
            _money_cell(ws, r, 9, rec["client_debt"])
        for col in range(1, 10):
            ws.cell(row=r, column=col).fill = _CHEK_FILL
            ws.cell(row=r, column=col).font = Font(bold=True)
        jami += float(rec.get("payable") or 0)
        qarzga += float(rec.get("debt") or 0)
        r += 1

        for it in rec.get("items") or []:
            ws.cell(row=r, column=4, value=it.get("name") or "")
            ws.cell(row=r, column=5, value=it.get("qty") or 0)
            _money_cell(ws, r, 6, it.get("price"))
            _money_cell(ws, r, 7, it.get("amount"))
            # Guruhlash: chek qatori ostidagi tovarlar yig'iladi/ochiladi.
            ws.row_dimensions[r].outline_level = 1
            ws.row_dimensions[r].hidden = False
            r += 1

    ws.cell(row=r + 1, column=4, value="JAMI").font = Font(bold=True)
    _money_cell(ws, r + 1, 7, jami).font = Font(bold=True)
    _money_cell(ws, r + 1, 8, qarzga).font = Font(bold=True)

    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(_OUT_DIR, f"cheklar-{_safe(title)}-{stamp}.xlsx")
    wb.save(path)
    return path


def _ledger_canvas(name: str, rows, debt_now: float, limit: int = 40):
    """Hisob-kitob rasmini chizadi → (rasm, kesish mumkin y nuqtalari).

    Rasmni PDF qilish uchun qatorlarni YARMIDAN kesmaslik kerak —
    shuning uchun har qator oxiridagi y qaytariladi.

    NEGA KERAK: Android telefonlarda .xlsx ni ochadigan dastur ko'pincha
    yo'q — fayl «PDF» deb ko'rsatiladi yoki umuman ochilmaydi. Rasm esa
    har qanday telefonda darrov ochiladi, suratga o'xshab galereyaga
    tushadi va uni boshqaga yuborish ham oson.

    Oxirgi `limit` ta qator chiziladi — rasm cheksiz uzun bo'lmasligi
    kerak; to'liq tarix Excel'da qoladi.
    """
    from PIL import Image, ImageDraw

    rows = list(rows or [])[-limit:]
    W = 1000
    pad, head_h, row_h, foot_h = 28, 190, 54, 96
    H = head_h + row_h * max(1, len(rows)) + foot_h

    YASHIL = (31, 111, 67)
    OCH = (240, 248, 242)
    MATN = (24, 32, 27)
    KUL = (120, 132, 124)
    QIZIL = (200, 50, 50)
    KOK = (28, 120, 72)

    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    breaks = []

    f_big = _fnt(44, True)
    f_h = _fnt(26, True)
    f = _fnt(24)
    f_s = _fnt(20)

    # Sarlavha — do'kon rangida
    d.rectangle([0, 0, W, head_h], fill=YASHIL)
    d.text((pad, 26), "CHINOR", font=f_big, fill=(255, 255, 255))
    d.text((pad, 80), f"Hisob-kitob · {name}", font=f_h, fill=(226, 240, 230))
    qarz = f"{round(debt_now):,}".replace(",", " ")
    holat = f"Joriy qarz: {qarz} so'm" if debt_now > 0 else "Qarz yo'q"
    d.text((pad, 128), holat, font=f_h, fill=(255, 255, 255))

    TUR = {"receipt": "Xarid", "payment": "To'lov", "manual": "Qo'lda",
           "fix": "Tuzatish", "opening": "Boshlang'ich"}
    CHEK = {"sale": "Xarid", "sale_deferred": "Nasiya", "return": "Qaytarish"}

    y = head_h
    for i, r in enumerate(rows):
        if i % 2 == 0:
            d.rectangle([0, y, W, y + row_h], fill=OCH)
        tur = (CHEK.get(r.get("kind"), "Xarid") if r.get("type") == "receipt"
               else TUR.get(r.get("type"), str(r.get("type") or "")))
        d.text((pad, y + 8), (r.get("at") or "")[:16], font=f_s, fill=KUL)
        d.text((pad, y + 28), tur, font=f, fill=MATN)
        ozg = float(r.get("change") or 0)
        rang = QIZIL if ozg > 0 else KOK
        belgi = "+" if ozg > 0 else ""
        txt = belgi + f"{round(ozg):,}".replace(",", " ")
        d.text((W - pad - d.textlength(txt, font=f), y + 10), txt,
               font=f, fill=rang)
        bal = f"{round(float(r.get('balance') or 0)):,}".replace(",", " ")
        d.text((W - pad - d.textlength(bal, font=f_s), y + 34), bal,
               font=f_s, fill=KUL)
        y += row_h
        breaks.append(y)

    d.rectangle([0, H - foot_h, W, H], fill=OCH)
    d.text((pad, H - foot_h + 18),
           f"Jami {len(rows)} ta harakat · {datetime.now():%d.%m.%Y %H:%M}",
           font=f_s, fill=KUL)
    d.text((pad, H - foot_h + 48),
           "Sergeli, Chinor bekati yoni · har kuni 8:00–22:00",
           font=f_s, fill=KUL)
    return im, breaks


def ledger_png(name: str, rows, debt_now: float, path: str = "",
               limit: int = 40) -> str:
    """Hisob-kitob — bitta uzun rasm (telefon uchun eng qulayi)."""
    im, _breaks = _ledger_canvas(name, rows, debt_now, limit)
    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(_OUT_DIR, f"hisob-{_safe(name)}-{stamp}.png")
    im.save(path, "PNG")
    return path


def ledger_pdf(name: str, rows, debt_now: float, path: str = "",
               limit: int = 120) -> str:
    """Hisob-kitob — PDF (A4 sahifalarga bo'lingan)."""
    im, breaks = _ledger_canvas(name, rows, debt_now, limit)
    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(_OUT_DIR, f"hisob-{_safe(name)}-{stamp}.pdf")
    return _save_pdf(im, breaks, path)


def _receipts_canvas(title: str, recs):
    """Cheklar rasmini chizadi → (rasm, kesish mumkin y nuqtalari).

    Mijoz «2-3 chekni yuboring» desa, har biriga alohida fayl yuborish
    noqulay: telefonda ularni yonma-yon ko'rib bo'lmaydi. Shuning uchun
    tanlangan cheklar bitta uzun rasmga, chek-chek ajratib chiziladi.

    `recs` — `_collect()` qaytargan shakl: {at, doc_no, kind, payable,
    debt, client_debt, items[{name, qty, price, amount}]}.
    """
    from PIL import Image, ImageDraw

    recs = list(recs or [])
    pad, W = 28, 1000
    head_h, chek_h, qator_h, jami_h, foot_h = 150, 72, 56, 64, 76
    H = head_h + foot_h + sum(
        chek_h + qator_h * len(r.get("items") or []) + jami_h for r in recs)
    H = max(H, head_h + foot_h + 80)

    YASHIL, OCH = (31, 111, 67), (240, 248, 242)
    MATN, KUL, QIZIL = (24, 32, 27), (120, 132, 124), (200, 50, 50)

    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    breaks = []
    f_big, f_h, f, f_s = _fnt(44, True), _fnt(26, True), _fnt(24), _fnt(20)

    d.rectangle([0, 0, W, head_h], fill=YASHIL)
    d.text((pad, 26), "CHINOR", font=f_big, fill=(255, 255, 255))
    d.text((pad, 80), f"Xaridlaringiz · {title}", font=f_h, fill=(226, 240, 230))
    d.text((pad, 114), f"{len(recs)} ta chek", font=f_s, fill=(200, 226, 208))

    y = head_h
    jami_hammasi = 0.0
    for r in recs:
        d.rectangle([0, y, W, y + chek_h], fill=OCH)
        bosh = f"{r.get('kind', 'savdo').capitalize()}"
        if r.get("doc_no"):
            bosh += f"  №{r['doc_no']}"
        d.text((pad, y + 10), bosh, font=f_h, fill=YASHIL)
        d.text((pad, y + 42), (r.get("at") or "")[:16], font=f_s, fill=KUL)
        y += chek_h

        for it in (r.get("items") or []):
            nom = str(it.get("name") or "")
            # Uzun nom ustunga sig'masin — kesib, uch nuqta qo'yamiz.
            while d.textlength(nom, font=f) > W - 2 * pad - 260 and len(nom) > 4:
                nom = nom[:-2]
            if nom != str(it.get("name") or ""):
                nom += "…"
            d.text((pad + 14, y + 10), nom, font=f, fill=MATN)
            miqdor = f"{float(it.get('qty') or 0):g} × {round(float(it.get('price') or 0)):,}"
            d.text((pad + 14, y + 34), miqdor.replace(",", " "), font=f_s, fill=KUL)
            summa = f"{round(float(it.get('amount') or 0)):,}".replace(",", " ")
            d.text((W - pad - d.textlength(summa, font=f), y + 16), summa,
                   font=f, fill=MATN)
            y += qator_h

        d.line([pad, y + 14, W - pad, y + 14], fill=(214, 228, 218), width=2)
        tolov = float(r.get("payable") or 0)
        jami_hammasi += tolov
        txt = f"Jami: {round(tolov):,}".replace(",", " ") + " so'm"
        d.text((W - pad - d.textlength(txt, font=f_h), y + 26), txt,
               font=f_h, fill=MATN)
        if float(r.get("debt") or 0) > 0:
            q = f"qarzga {round(float(r['debt'])):,}".replace(",", " ")
            d.text((pad + 14, y + 30), q, font=f_s, fill=QIZIL)
        y += jami_h
        # Chekning oxiri — PDF sahifasini shu yerda kesish mumkin.
        breaks.append(y)

    d.rectangle([0, H - foot_h, W, H], fill=OCH)
    jt = f"HAMMASI: {round(jami_hammasi):,}".replace(",", " ") + " so'm"
    d.text((pad, H - foot_h + 14), jt, font=f_h, fill=YASHIL)
    d.text((pad, H - foot_h + 46),
           "Sergeli, Chinor bekati yoni · har kuni 8:00–22:00",
           font=f_s, fill=KUL)
    return im, breaks


def receipts_png(title: str, recs, path: str = "") -> str:
    """Tanlangan cheklar — bitta uzun rasm."""
    im, _breaks = _receipts_canvas(title, recs)
    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(_OUT_DIR, f"cheklar-{_safe(title)}-{stamp}.png")
    im.save(path, "PNG")
    return path


def receipts_pdf(title: str, recs, path: str = "") -> str:
    """Tanlangan cheklar — PDF (chop etishga va arxivga qulay).

    NEGA PDF HAM KERAK. Rasmni telefon darrov ko'rsatadi, lekin uni
    chop etish yoki hujjatga ilova qilish noqulay; Excel esa Android'da
    ko'pincha umuman ochilmaydi. PDF ikkalasining o'rtasi: har qanday
    qurilmada ochiladi va chop etilsa chek ko'rinishini saqlaydi.
    """
    im, breaks = _receipts_canvas(title, recs)
    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(_OUT_DIR, f"cheklar-{_safe(title)}-{stamp}.pdf")
    return _save_pdf(im, breaks, path)


def photo_yaroqli(path: str) -> bool:
    """Shu rasmni Telegram «photo» sifatida qabul qiladimi?

    Telegram photo uchun chegara qo'yadi: tomonlar yig'indisi 10000 px dan
    va fayl 10 MB dan oshmasligi kerak. 30 ta chek chizilgan rasm esa
    osongina 12 000 px bo'lib ketadi — o'shanda `sendPhoto` «PHOTO_INVALID
    _DIMENSIONS» deb rad etadi va mijozga HECH NARSA bormaydi. Shuning
    uchun chaqiruvchi avval shuni so'raydi: yaroqsiz bo'lsa HUJJAT qilib
    yuboradi (sifati ham saqlanadi).
    """
    try:
        from PIL import Image

        if os.path.getsize(path) > 9 * 1024 * 1024:
            return False
        with Image.open(path) as im:
            w, h = im.size
        return (w + h) <= 9000 and h <= 8000
    except Exception:
        return True


# ── Uzun rasmni A4 sahifalarga bo'lish ──────────────────────────────────────
# Rasm bitta ustun bo'lib cho'ziladi (eni 1000 px). A4 ning eni 210 mm,
# bo'yi 297 mm — demak 1000 px eni uchun bir sahifa ≈ 1414 px bo'ladi.
# Kesish joyi QATOR OXIRIGA to'g'rilanadi: aks holda chek nomi yoki
# summasi ikki sahifaga bo'linib ketardi.

_PAGE_H = 1414


def _save_pdf(im, breaks, path: str) -> str:
    pages = _paginate(im, breaks)
    first, rest = pages[0], pages[1:]
    first.save(path, "PDF", resolution=120.0, save_all=bool(rest),
               append_images=rest)
    return path


def _paginate(im, breaks, page_h: int = _PAGE_H):
    from PIL import Image

    W, H = im.size
    nuqta = sorted({int(b) for b in (breaks or []) if 0 < int(b) < H}) + [H]
    sahifalar, tep = [], 0
    while tep < H:
        mumkin = [b for b in nuqta if tep < b <= tep + page_h]
        past = max(mumkin) if mumkin else min(H, tep + page_h)
        bolak = im.crop((0, tep, W, past))
        sahifa = Image.new("RGB", (W, page_h), (255, 255, 255))
        sahifa.paste(bolak, (0, 0))
        sahifalar.append(sahifa)
        tep = past
    return sahifalar or [im]


def _fnt(size: int, bold: bool = False):
    """Shrift — serverda faqat DejaVu bor, topilmasa standart."""
    from PIL import ImageFont

    base = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "..", "assets", "fonts")
    nom = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    try:
        return ImageFont.truetype(os.path.normpath(os.path.join(base, nom)), size)
    except Exception:
        try:
            return ImageFont.truetype(nom, size)
        except Exception:
            return ImageFont.load_default()


def ledger_xlsx(name: str, rows, debt_now: float, path: str = "") -> str:
    """Mijozning hisob-kitob daftari — bitta varaq."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisob-kitob"
    heads = ["Sana", "Turi", "Hujjat", "Izoh", "O'zgarish", "Qoldiq"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font = _HEAD_FILL, _HEAD_FONT
    ws.freeze_panes = "A2"
    for i, w in enumerate([19, 12, 10, 44, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    _TUR = {"payment": "To'lov", "manual": "Qo'lda",
            "fix": "Tuzatish", "opening": "Boshlang'ich"}
    # Chek turini ham ajratamiz: qaytarish qarzni KAMAYTIRADI, uni oddiy
    # xarid deb ko'rsatish daftarni tushunarsiz qilardi.
    _CHEK = {"sale": "Xarid", "sale_deferred": "Nasiya", "return": "Qaytarish"}
    r = 2
    for row in rows:
        ws.cell(row=r, column=1, value=row.get("at") or "")
        tur = (_CHEK.get(row.get("kind"), "Xarid") if row.get("type") == "receipt"
               else _TUR.get(row.get("type"), row.get("type")))
        ws.cell(row=r, column=2, value=tur)
        ws.cell(row=r, column=3, value=row.get("doc_no") or "")
        ws.cell(row=r, column=4, value=row.get("note") or "")
        _money_cell(ws, r, 5, row.get("change"))
        _money_cell(ws, r, 6, row.get("balance"))
        r += 1

    ws.cell(row=r + 1, column=4, value="JORIY QARZ").font = Font(bold=True)
    _money_cell(ws, r + 1, 6, debt_now).font = Font(bold=True)

    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(_OUT_DIR, f"hisob-{_safe(name)}-{stamp}.xlsx")
    wb.save(path)
    return path
